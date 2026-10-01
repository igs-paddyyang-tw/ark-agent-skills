#!/usr/bin/env python3
"""qt_probtable — 機率規格書 View 軌（xlsx 公版）：由 prob-spec.meta.json 指到的來源（odds.json / qt-config / rtp-report / gdd / decisions）
直接 dump 成 `data/prob/<slug>/prob-spec.xlsx`，分頁照「機率規格書共用框架版範本」：
製作方針 / 規格簡述 / 數據資料 / 機率流程圖 / 參數表 / Main Game Strip / Free Game Strip / 轉置 Strip / 隱性規則 / _meta。

黃金律（game-prob-table-excel）：每個數字從來源 dump，不手抄；派生格一律活公式（COUNTIF / SUM）。
null → 「待決議」黃底；每個值旁有來源欄（qt-config.provenance）。
用法: python qt_probtable.py --dir data/prob/<slug>   （先跑 qt_probspec）
"""
from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qt_common as C  # noqa: E402

PENDING = "待決議"
GUIDE = ["機率規格書製作方針", "", "1.  規格簡述", "機率人員對該玩法之簡述(競品分析模式)", "特色觸發 & 進行方式", "2.  數據資料", "定義必要數據，以該數據製作框架",
         "3.  機率流程圖", "需於機率文件附上可供編輯之流程圖", "4.  參數表", "加入機率隱性規則/「更改需與編導規格書同步」(不強制放同一頁)", "隱性規則須加上設計目的",
         "5.  競品資料", "(選用)加入收錄目的"]
SCALARS = [("extra_rate", "額外押注倍率"), ("enter_free_game_gate", "進免費遊戲門檻"), ("free_game_num", "免費遊戲手數"), ("free_game_retrigger_gate", "retrigger 門檻"),
           ("free_game_retrigger_spin_num", "retrigger 加手數"), ("max_win", "MaxWin 倍數"), ("is_item_card", "道具卡")]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True, help="data/prob/<slug>（需有 prob-spec.meta.json）")
    a = ap.parse_args()
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter as col
    except ImportError:
        C.fail("MISSING_DEP", "需要 openpyxl", "pip install openpyxl")
    d = pathlib.Path(a.dir)
    meta_p = d / "prob-spec.meta.json"
    if not meta_p.exists():
        C.fail("BAD_INPUT", f"缺 {meta_p}", "先跑 qt_probspec.py")
    meta = json.loads(meta_p.read_text(encoding="utf-8"))
    cwd = pathlib.Path.cwd()

    def locate(rel: str) -> pathlib.Path | None:
        q = pathlib.Path(rel)
        if q.is_absolute():
            return q if q.exists() else None
        for base in [cwd, *d.resolve().parents]:
            if (base / q).exists():
                return base / q
        return None

    S = {k: locate(v["path"]) for k, v in (meta.get("sources") or {}).items()}
    for k in ("qt-config", "odds"):
        if not S.get(k):
            C.fail("BAD_INPUT", f"來源 {k} 找不到：{meta['sources'][k]['path']}")
    cfg = C.yaml_load(S["qt-config"])
    odds = json.loads(S["odds"].read_text(encoding="utf-8"))
    eo = odds.get("extra_odds") or {}
    prov = cfg.get("provenance") or {}
    st = cfg.get("structure") or {}
    reels = int(st.get("reels") or 5)
    symbols = st.get("symbols") or []
    code_of = {str(s.get("sym_id")): str(s.get("code")) for s in symbols}
    rep = json.loads(S["rtp-report"].read_text(encoding="utf-8")) if S.get("rtp-report") else None
    gdd = C.yaml_load(S["gdd"]) if S.get("gdd") else None
    dec = C.yaml_load(S["decisions"]) if S.get("decisions") else None
    sm = S["state-machine"].read_text(encoding="utf-8") if S.get("state-machine") else None

    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="5B4636")
    TF = Font(bold=True); PF = PatternFill("solid", fgColor="FFE699"); SF = PatternFill("solid", fgColor="F2F2F2")
    wb = openpyxl.Workbook()
    n_pending = 0

    def hdr(ws, r, c, vals):
        for i, v in enumerate(vals):
            x = ws.cell(r, c + i, v); x.font = H; x.fill = HF; x.alignment = Alignment(horizontal="center")

    def val(ws, r, c, v, src=None):
        nonlocal n_pending
        if v is None:
            x = ws.cell(r, c, PENDING); x.fill = PF; n_pending += 1
        else:
            x = ws.cell(r, c, v)
        if src is not None:
            s = ws.cell(r, c + 1, src or "—"); s.font = Font(color="7F7F7F", italic=True)
        return x

    # 1 製作方針
    ws = wb.active; ws.title = "機率規格書製作方針"
    for i, t in enumerate(GUIDE, 1):
        x = ws.cell(i, 1, t); x.font = TF if re.match(r"^\d\.", t) or i == 1 else Font()
    ws.column_dimensions["A"].width = 70

    # 2 規格簡述
    ws = wb.create_sheet("規格簡述"); r = 1
    if gdd:
        hdr(ws, r, 1, ["項目", "內容", "來源"]); r += 1
        for kv in gdd.get("spec") or []:
            if kv.get("kv") is False:
                continue
            ws.cell(r, 1, str(kv.get("k"))); ws.cell(r, 2, str(kv.get("v"))); ws.cell(r, 3, "SPEC gdd.yaml.spec"); r += 1
        r += 1
        for f in gdd.get("features") or []:
            x = ws.cell(r, 1, f.get("title") or f.get("id")); x.font = TF; r += 1
            rp = S["gdd"].parent / (f.get("rules") or "")
            if f.get("rules") and rp.exists():
                for ln in rp.read_text(encoding="utf-8").splitlines():
                    s = ln.strip()
                    if not s or s.startswith("<!--"):
                        continue
                    ws.cell(r, 2, re.sub(r"\*\*(.+?)\*\*", r"\1", s)); ws.cell(r, 3, "SPEC " + f["rules"]); r += 1
            r += 1
    else:
        ws.cell(r, 1, "無 gdd-pack：規格簡述待補")
    r += 1; x = ws.cell(r, 1, "機率人員簡述（競品分析模式）"); x.font = TF; ws.cell(r + 1, 2, "（填寫處；Content 軌 prob-spec.md 的 human:summary 區為正本）")
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 90; ws.column_dimensions["C"].width = 24

    # 3 數據資料（公版座標：G2 表頭、F3.. 階段）
    ws = wb.create_sheet("數據資料")
    hdr(ws, 2, 7, ["RTP", "Hit %", "觸發率", "平均觸發局數", "倍率", "最大倍率", "平均局數"])
    ws.cell(3, 3, "快測" + (f"（{rep.get('version')}，{int(rep.get('spins') or 0):,} 手）" if rep else "（未快測）"))
    ws.cell(5, 3, f"來源：qt-report/rtp-report.json；verdict {meta.get('verdict')}" if rep else "跑 qt_run 後重產：qt_probspec → qt_probtable")
    games = list((rep.get("games") or {}).items()) if rep else [("Main Game", {}), ("Feature Game", {}), ("Jackpot", {}), ("Total", {})]
    for i, (name, g) in enumerate(games):
        rr = 3 + i; ws.cell(rr, 6, name).font = TF
        freq = g.get("freq")
        cells = [g.get("rtp"), g.get("hitrate"), (1.0 / freq) if freq else None, freq, g.get("multi"), g.get("maxmulti"), g.get("playtimes")]
        for j, v in enumerate(cells):
            x = ws.cell(rr, 7 + j, v if v is not None else "--")
            if j in (0, 1, 2) and v is not None:
                x.number_format = "0.00%"
    if rep:
        rr = 3 + len(games) + 1
        for k, lab in (("sd", "SD"), ("pi", "P.I."), ("pay_out_rate", "Pay Out Rate")):
            ws.cell(rr, 6, lab); ws.cell(rr, 7, (rep.get("detail") or {}).get(k)); rr += 1
        for c in (rep.get("gate") or {}).get("checks") or []:
            ws.cell(rr, 6, c["id"]); ws.cell(rr, 7, c["status"]); ws.cell(rr, 8, c["msg"]); rr += 1
    for c_ in range(6, 14):
        ws.column_dimensions[col(c_)].width = 16

    # 4 機率流程圖
    ws = wb.create_sheet("機率流程圖")
    ws.cell(1, 1, "附可供編輯之流程圖檔於機率文件資料夾內（drawio 為 View 軌下一版；以下為 dev-spec/state-machine.md 的 mermaid 原文）").font = TF
    if sm:
        m = re.search(r"```mermaid\n(.*?)```", sm, re.S)
        for i, ln in enumerate((m.group(1) if m else sm).splitlines(), 3):
            ws.cell(i, 1, ln)
    else:
        ws.cell(3, 1, "無 state-machine.md；待 gs_run --stage dev")
    ws.column_dimensions["A"].width = 80

    # 5 參數表
    ws = wb.create_sheet("參數表")
    ws.cell(2, 1, "基本資訊").font = TF
    ws.cell(2, 3, "盤面"); val(ws, 2, 4, f"{st.get('reels')}x{st.get('rows')}" if st.get("reels") and st.get("rows") else None, prov.get("structure.reels"))
    rr = 3
    for kv in (gdd or {}).get("spec") or []:
        if any(t in str(kv.get("k")) for t in ("對獎", "線", "收費", "成本")):
            ws.cell(rr, 3, str(kv.get("k"))); ws.cell(rr, 4, str(kv.get("v"))); ws.cell(rr, 5, "SPEC gdd.yaml.spec"); rr += 1
    rr = max(rr + 1, 7)
    ws.cell(rr, 1, "Odds表").font = TF; ws.cell(rr, 3, "Pay table"); rr += 1
    hdr(ws, rr, 3, ["Description", "Symbol", "sym_id"] + [str(k) for k in range(1, reels + 1)] + ["來源"]); rr += 1
    odds_first = rr
    for s in symbols:
        sid = str(s.get("sym_id")); arr = (eo.get("odds") or {}).get(sid)
        ws.cell(rr, 3, s.get("group") or ""); ws.cell(rr, 4, s.get("code")); ws.cell(rr, 5, s.get("sym_id"))
        for k in range(reels):
            val(ws, rr, 6 + k, (arr[k] if arr is not None and k < len(arr) else (None if arr is None else 0)))
        src = prov.get(f"extra_odds.odds.{sid}") or prov.get("extra_odds.odds")
        ws.cell(rr, 6 + reels, src or ("—" if arr is None else "UNKNOWN")).font = Font(color="7F7F7F", italic=True); rr += 1
    rr += 2
    ws.cell(rr, 1, "表P-1").font = TF; ws.cell(rr, 3, "觸發與上限參數"); rr += 1
    hdr(ws, rr, 3, ["參數", "值", "來源"]); rr += 1
    for key, lab in SCALARS:
        ws.cell(rr, 3, lab); val(ws, rr, 4, eo.get(key), prov.get(f"extra_odds.{key}")); rr += 1
    for key, lab in (("jackpot_multiple", "Jackpot 基底倍數"), ("jackpot_acc_rate", "Jackpot 累積率")):
        v = eo.get(key); ws.cell(rr, 3, lab); val(ws, rr, 4, ", ".join(map(str, v)) if isinstance(v, list) else None, prov.get(f"extra_odds.{key}")); rr += 1
    rr += 2

    def index_table(label, title, iset, wt, key):
        nonlocal rr
        ws.cell(rr, 1, label).font = TF; ws.cell(rr, 3, title); rr += 1
        if not isinstance(iset, list) or not iset:
            val(ws, rr, 3, None, prov.get(key)); rr += 2; return
        n = max(len(x) for x in iset if isinstance(x, list))
        two_d = isinstance(wt, list) and wt and isinstance(wt[0], list)
        wcols = ["Weight(超強)", "Weight(強)", "Weight(中)", "Weight(弱)"] if two_d else ["Weight"]
        hdr(ws, rr, 3, ["Index"] + [f"R{i + 1}" for i in range(n)] + wcols + ["來源"]); rr += 1
        first = rr
        for i, combo in enumerate(iset):
            ws.cell(rr, 3, i)
            for j, g in enumerate(combo):
                ws.cell(rr, 4 + j, g)
            if two_d:
                for t in range(4):
                    val(ws, rr, 4 + n + t, wt[t][i] if t < len(wt) and i < len(wt[t]) else None)
            else:
                val(ws, rr, 4 + n, wt[i] if isinstance(wt, list) and i < len(wt) else None)
            ws.cell(rr, 4 + n + len(wcols), prov.get(key.replace("index_set", "set_weight")) or "—").font = Font(color="7F7F7F", italic=True)
            rr += 1
        for t in range(len(wcols)):
            c_ = col(4 + n + t); ws.cell(rr, 4 + n + t, f"=SUM({c_}{first}:{c_}{rr - 1})").font = TF
        rr += 2

    index_table("表M-1", "Main Game 轉輪帶設定（index_set × weight，K-022）", eo.get("main_game_reel_index_set"), eo.get("main_game_reel_set_weight"), "extra_odds.main_game_reel_index_set")
    ws.cell(rr, 1, "表F-1").font = TF; ws.cell(rr, 3, "決定強中弱組"); rr += 1
    hdr(ws, rr, 3, ["組別", "Weight", "來源"]); rr += 1
    ftw = eo.get("free_game_type_weight"); first = rr
    for i, lab in enumerate(["超強", "強", "中", "弱"]):
        ws.cell(rr, 3, lab); val(ws, rr, 4, ftw[i] if isinstance(ftw, list) and i < len(ftw) else None, prov.get("extra_odds.free_game_type_weight")); rr += 1
    ws.cell(rr, 4, f"=SUM(D{first}:D{rr - 1})").font = TF; rr += 2
    index_table("表F-2", "Free Game 轉輪帶設定（四型各自權重）", eo.get("free_game_reel_index_set"), eo.get("free_game_reel_set_weight"), "extra_odds.free_game_reel_index_set")
    for c_ in range(1, 16):
        ws.column_dimensions[col(c_)].width = 14
    ws.column_dimensions["C"].width = 26

    # 6/7 Strips（活公式 COUNTIF）
    def strip_sheet(name, data):
        ws = wb.create_sheet(name)
        if not isinstance(data, list) or not data:
            ws.cell(1, 1, f"{PENDING}（來源為 null）").fill = PF; return 0
        r0 = 1; total = 0
        for gi, grp in enumerate(data):
            n = len(grp or []); L = max((len(x or []) for x in grp or []), default=0)
            ws.cell(r0 + 1, 2, f"Reels_{gi}").font = TF
            hdr(ws, r0 + 1, 3, [f"R{i + 1}" for i in range(n)])
            for ri, reel in enumerate(grp or []):
                for k, v in enumerate(reel or []):
                    ws.cell(r0 + 2 + k, 1 + 1, k + 1) if ri == 0 else None
                    ws.cell(r0 + 2 + k, 3 + ri, code_of.get(str(v), str(v))); total += 1
            cc = 3 + n + 1
            hdr(ws, r0 + 1, cc, ["總顆數"] + [f"R{i + 1}" for i in range(n)])
            last = r0 + 1 + L
            for si, s in enumerate(symbols):
                ws.cell(r0 + 2 + si, cc, s.get("code"))
                for ri in range(n):
                    c_ = col(3 + ri)
                    ws.cell(r0 + 2 + si, cc + 1 + ri, f"=COUNTIF({c_}${r0 + 2}:{c_}${last},${col(cc)}{r0 + 2 + si})")
            tr = r0 + 2 + len(symbols)
            ws.cell(tr, cc, "Total").font = TF
            for ri in range(n):
                c_ = col(cc + 1 + ri); ws.cell(tr, cc + 1 + ri, f"=SUM({c_}{r0 + 2}:{c_}{tr - 1})").font = TF
            r0 = max(last, tr) + 2
        return total

    n_main = strip_sheet("Main Game Strip", odds.get("main_game_reel_data"))
    n_free = strip_sheet("Free Game Strip", odds.get("free_game_reel_data"))

    # 8 轉置 Strip（Symbol ↔ 轉換用 Code）
    ws = wb.create_sheet("轉置 Strip")
    hdr(ws, 1, 1, ["Symbol", "轉換用Code", "名稱", "組"])
    for i, s in enumerate(symbols, 2):
        ws.cell(i, 1, s.get("code")); ws.cell(i, 2, s.get("sym_id")); ws.cell(i, 3, s.get("name") or ""); ws.cell(i, 4, s.get("group") or "")

    # 9 隱性規則
    ws = wb.create_sheet("隱性規則")
    hdr(ws, 1, 1, ["編號", "情況", "說明", "設計目的", "來源"]); r = 2
    for dd in (dec or {}).get("decisions") or []:
        if dd.get("decision") == "defer" or dd.get("value") is None:
            continue
        ws.cell(r, 1, dd.get("decision_id")); ws.cell(r, 2, dd.get("topic")); ws.cell(r, 3, json.dumps(dd.get("value"), ensure_ascii=False)); ws.cell(r, 4, dd.get("reason") or ""); ws.cell(r, 5, f"DECIDED {dd.get('decided_by') or ''}"); r += 1
    if r == 2:
        ws.cell(2, 2, "尚無已定案決議；更改需與編導規格書同步，每條須有設計目的")
    for c_, w in zip("ABCDE", (10, 28, 40, 40, 24)):
        ws.column_dimensions[c_].width = w

    # _meta（雙軌戳記）
    ws = wb.create_sheet("_meta"); ws.sheet_state = "hidden"
    ws.cell(1, 1, "content-src"); ws.cell(1, 2, "prob-spec.md"); ws.cell(1, 3, C.sha16(d / "prob-spec.md") if (d / "prob-spec.md").exists() else "")
    ws.cell(2, 1, "odds"); ws.cell(2, 2, meta["sources"]["odds"]["path"]); ws.cell(2, 3, C.sha16(S["odds"]))
    ws.cell(3, 1, "generated_by"); ws.cell(3, 2, "ark-game-quicktest qt_probtable"); ws.cell(3, 3, meta.get("generated"))

    out = d / "prob-spec.xlsx"
    wb.save(out)
    C.emit({"xlsx": out.as_posix(), "sheets": wb.sheetnames, "pending_cells": n_pending, "strip_cells": {"main": n_main, "free": n_free},
            "delivery": f"prob-spec.xlsx 已落盤 {out.as_posix()}｜{len(wb.sheetnames)} 分頁｜待決議格 {n_pending}｜輪帶格 main {n_main} / free {n_free}（COUNTIF 活公式）"}, {"stage": "probtable"})


if __name__ == "__main__":
    main()
