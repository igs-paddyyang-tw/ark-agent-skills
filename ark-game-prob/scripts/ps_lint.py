#!/usr/bin/env python3
"""ps_lint — 機率規格書 Content 軌守門（data/prob/<slug>/prob-spec.md）。

規則（deterministic）：
  PS-SECTIONS  六段齊全、順序固定（1 規格簡述 … 6 隱性規則）
  PS-NUM       全文每個數字（人工區除外）必須能回溯：來源檔全文（meta.sources）∪ 派生值（meta.derived，含公式來源）
  PS-PROV      §4 參數表每個非「待決議」的值都要有來源（provenance）；來源為 — / UNKNOWN → error
  PS-STALE     meta.sources 的 sha 與現況不一致 → error（來源更新後未重產）；html / xlsx(_meta) 戳記 ≠ md sha → error
  PS-HTML-EMPTY html 疑似空殼：去標籤純文字命中六段標題 < 4 → error；< 6 → warn（過戳記 ≠ 內容完整）
  PS-HUMAN     人工區含數字 → warn（未受 PS-NUM 保護，請自行標來源）
  PS-INJECT    指令覆寫句型 / 零寬字元 → error
用法: python ps_lint.py --dir data/prob/<slug>   → lint-report.json；error exit 3
"""
from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ps_common as C  # noqa: E402

PENDING = "待決議"
SECTIONS = ["1. 規格簡述", "2. 數據資料", "3. 機率流程圖", "4. 參數表", "5. 競品資料／手法對應", "6. 隱性規則"]
HUMAN_RE = re.compile(r"<!--\s*human:([a-z0-9_-]+)\s*-->\n?(.*?)<!--\s*/human\s*-->", re.S)
NUM_RE = re.compile(r"(?<![A-Za-z_\-\d.#/])\d[\d,]*(?:\.\d+)?")
INJECT_RE = re.compile(r"(請忽略(以上|之前|先前)|ignore (all )?(previous|above|prior) instructions|you are now|系統提示|system prompt)", re.I)
ZW_RE = re.compile(r"[​-‏  ﻿]")


def norm(tok: str) -> str:
    return tok.replace(",", "").rstrip("%")


HEX_RE = re.compile(r"(?<![0-9a-zA-Z])[0-9a-f]{12,64}(?![0-9a-zA-Z])")


def numbers_in(text: str) -> set[str]:
    text = HEX_RE.sub(" ", text)  # sha / id 不是數字
    return {norm(m.group(0)) for m in NUM_RE.finditer(text)}


BINARY_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".zip", ".bin", ".png", ".jpg", ".exe"}
JSON_SKIP_KEYS = {"r", "c0", "col", "anchor", "generated", "dims", "sha256_16", "path"}


def json_text(o) -> str:
    """JSON 值攤平成文字：數字用 repr、字串原文（含真正換行），供 numbers_in 建允許集。"""
    out: list[str] = []
    def walk(x):
        if isinstance(x, bool):
            return
        if isinstance(x, (int, float)):
            out.append(repr(x) if isinstance(x, float) else str(x))
        elif isinstance(x, str):
            out.append(x)
        elif isinstance(x, dict):
            for k, v in x.items():
                if k in JSON_SKIP_KEYS:  # 座標／時間戳不是規格數字
                    continue
                walk(v)
        elif isinstance(x, (list, tuple)):
            for v in x:
                walk(v)
    walk(o)
    return "\n".join(out)


def split_fm(text: str) -> tuple[str, str]:
    if text.startswith("---\n"):
        end = text.find("\n---\n", 4)
        if end > 0:
            return text[4:end], text[end + 5:]
    return "", text


def lint(d: pathlib.Path) -> dict:
    V: list[dict] = []

    def add(rule, where, msg, sev="error"):
        V.append({"rule": rule, "severity": sev, "where": where, "message": msg})

    md_p, meta_p, html_p = d / "prob-spec.md", d / "prob-spec.meta.json", d / "prob-spec.html"
    if not md_p.exists() or not meta_p.exists():
        add("PS-STALE", d.as_posix(), "缺 prob-spec.md 或 prob-spec.meta.json")
        return {"summary": {"errors": 1, "warnings": 0}, "violations": V}
    text = md_p.read_text(encoding="utf-8")
    meta = json.loads(meta_p.read_text(encoding="utf-8"))
    fm, body = split_fm(text)

    # PS-INJECT
    for i, ln in enumerate(text.splitlines(), 1):
        if INJECT_RE.search(ln):
            add("PS-INJECT", f"prob-spec.md:{i}", "指令覆寫句型")
        if ZW_RE.search(ln):
            add("PS-INJECT", f"prob-spec.md:{i}", "零寬 / 方向控制字元")

    # PS-SECTIONS
    heads = [ln[3:].strip() for ln in body.splitlines() if ln.startswith("## ")]
    core = [h for h in heads if re.match(r"^\d\. ", h)]
    if core != SECTIONS:
        add("PS-SECTIONS", "prob-spec.md", f"章節應為 {SECTIONS}，實際 {core}")

    # PS-STALE
    allowed = set()
    def locate(rel: str) -> pathlib.Path:
        q = pathlib.Path(rel)
        if q.is_absolute():
            return q
        for base in [pathlib.Path.cwd(), *d.resolve().parents]:  # 相對專案根；lint 可從任何 cwd 跑
            if (base / q).exists():
                return base / q
        return q

    for k, s in (meta.get("sources") or {}).items():
        p = locate(s["path"])
        if not p.exists():
            add("PS-STALE", k, f"來源不存在：{p}")
            continue
        if C.sha16(p) != s.get("sha256_16"):
            add("PS-STALE", k, f"來源已變更（{p.name}），請重產 prob-spec")
        try:
            if p.suffix.lower() in BINARY_SUFFIXES:  # 二進位來源只驗 sha，不拿位元組當數字
                pass
            elif p.suffix.lower() == ".json":  # 結構化來源：走 JSON 值，不受 \n 逸出字元影響
                allowed |= numbers_in(json_text(json.loads(p.read_text(encoding="utf-8"))))
            else:
                allowed |= numbers_in(p.read_text(encoding="utf-8", errors="replace"))
        except Exception:  # noqa: BLE001
            pass
    if html_p.exists():
        html_text = html_p.read_text(encoding="utf-8", errors="replace")
        m = re.search(r"<!-- content-src: prob-spec\.md sha256:([0-9a-f]{16}) -->", html_text)
        if not m:
            add("PS-STALE", "prob-spec.html", "缺 content-src 戳記")
        elif m.group(1) != C.sha16(md_p):
            add("PS-STALE", "prob-spec.html", "html 戳記與 md 不符（md 改了未重渲染）")
        # PS-HTML-EMPTY：防空殼——html 須實際渲染六段內容，不能只有戳記 + 一句佔位。
        # 去標籤 / script / style 後的純文字，應命中六段標題中的多數；否則判為空殼。
        _stripped = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html_text)
        _plain = re.sub(r"(?s)<[^>]+>", " ", _stripped)
        _hit = sum(1 for s in SECTIONS if s.split(". ", 1)[-1] in _plain)
        if _hit < 4:
            add("PS-HTML-EMPTY", "prob-spec.html",
                f"html 疑似空殼：純文字只命中 {_hit}/6 段標題（過戳記 ≠ 內容完整）。"
                "build_html 須實際渲染六段內容（表格/清單/流程圖），非只放佔位句。")
        elif _hit < 6:
            add("PS-HTML-EMPTY", "prob-spec.html",
                f"html 內容不完整：純文字只命中 {_hit}/6 段標題，請確認六段都有渲染。", "warning")
    else:
        add("PS-STALE", "prob-spec.html", "缺 View 軌 html", "warning")
    xlsx_p = d / "prob-spec.xlsx"
    if xlsx_p.exists():
        try:
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(xlsx_p, read_only=True)
            if "_meta" not in wb.sheetnames:
                add("PS-STALE", "prob-spec.xlsx", "缺 _meta 戳記分頁")
            else:
                rows = {r[0]: (r[1], r[2]) for r in wb["_meta"].iter_rows(values_only=True) if r and r[0]}
                if rows.get("content-src", ("", ""))[1] != C.sha16(md_p):
                    add("PS-STALE", "prob-spec.xlsx", "xlsx 戳記與 md 不符（重跑 ps_probtable）")
                odds_src = (meta.get("sources") or {}).get("odds") or {}
                if odds_src and rows.get("odds", ("", ""))[1] != odds_src.get("sha256_16"):
                    add("PS-STALE", "prob-spec.xlsx", "xlsx 的 odds sha 與 meta 不符")
        except ImportError:
            add("PS-STALE", "prob-spec.xlsx", "無 openpyxl，略過 xlsx 戳記檢查", "warning")
    allowed |= {norm(str(x.get("value"))) for x in meta.get("derived") or []}
    allowed |= {str(meta.get("pending_params")), "0", "1"}

    # 人工區：剝掉再驗數字
    human_blocks = HUMAN_RE.findall(body)
    for hid, content in human_blocks:
        if numbers_in(content):
            add("PS-HUMAN", f"human:{hid}", "人工區含數字，不受 PS-NUM 保護，請在文中標來源", "warning")
    stripped = HUMAN_RE.sub("", body)

    # PS-NUM
    for i, ln in enumerate(stripped.splitlines(), 1):
        if ln.startswith("#") or ln.startswith("```") or ln.startswith("<!--"):
            continue
        for tok in numbers_in(ln) - allowed:
            add("PS-NUM", f"prob-spec.md 行「{ln[:40]}」", f"數字 '{tok}' 無法回溯到來源 / 派生值")

    # PS-PROV（§4）
    s4 = stripped.split("## 4. 參數表", 1)[1].split("## 5.", 1)[0] if "## 4. 參數表" in stripped else ""
    header: list[str] | None = None
    for ln in s4.splitlines():
        if not ln.startswith("|"):
            header = None
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        if header is None:
            header = cells
            continue
        if not header or header[-1] != "來源" or len(cells) != len(header):
            continue
        vcols = [i for i, h in enumerate(header) if re.fullmatch(r"x\d+", h)] or ([header.index("權重")] if "權重" in header else [1])
        vals = [cells[i] for i in vcols if i < len(cells)]
        src = cells[-1]
        has_value = any(v not in (PENDING, "—", "") for v in vals)
        if has_value and (src in ("—", "") or "UNKNOWN" in src):
            add("PS-PROV", f"§4 {cells[0]}", f"有值但無來源（{src or '空'}）：{vals}")
        if not has_value and any(v == PENDING for v in vals) and src not in ("—", "") and "UNKNOWN" not in src:
            add("PS-PROV", f"§4 {cells[0]}", f"標 {PENDING} 卻帶來源 {src}", "warning")

    errs = sum(1 for v in V if v["severity"] == "error")
    warns = len(V) - errs
    pend = stripped.count(PENDING)
    return {"summary": {"errors": errs, "warnings": warns, "pending_marks": pend, "sections": len(core), "sources": len(meta.get("sources") or {})}, "violations": V}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    a = ap.parse_args()
    d = pathlib.Path(a.dir)
    rep = lint(d)
    C.atomic_write(d / "lint-report.json", json.dumps(rep, ensure_ascii=False, indent=1))
    if rep["summary"]["errors"]:
        C.fail("GATE_BLOCKED", f"{rep['summary']['errors']} 個 lint error", "見 lint-report.json", rep)
    C.emit({"status": "PASS", **rep["summary"], "warnings_detail": [v for v in rep["violations"] if v["severity"] == "warning"],
            "delivery": f"ps_lint PASS｜{rep['summary']['sections']} 段｜{rep['summary']['pending_marks']} 處待決議｜warnings {rep['summary']['warnings']}"}, {"stage": "ps_lint"})


if __name__ == "__main__":
    main()
