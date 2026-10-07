"""ps_xlsx / _xlsx_lib / ps_config_load 守門測試。

反證原則：故意違規必紅、還原必綠。
- 文字型 '=' 開頭 → 強制字串（不被當公式）
- None 值 → 待決議黃底（block_table）；檢核表輸入格 → 留空淡藍
- 活公式字串以 '=' 開頭、data_type='f'
- config_load phase/shape/table 分類正確、缺值進 missing
- _meta 隱藏頁 + sha 戳記
需 openpyxl（缺則整檔 skip）。
"""
from __future__ import annotations

import json
import pathlib
import sys
import tempfile

import pytest

SCRIPTS = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SCRIPTS))

pytest.importorskip("openpyxl")
pytest.importorskip("yaml")

import _xlsx_lib as X  # noqa: E402
import ps_config_load as CL  # noqa: E402
import ps_xlsx as PX  # noqa: E402
import openpyxl  # noqa: E402


def _tmp(suffix=".xlsx"):
    return pathlib.Path(tempfile.mktemp(suffix=suffix))


# ── _xlsx_lib ──
def test_text_equals_forced_string():
    """文字型 '=規格' 不可被當公式（data_type='s'）。"""
    wb = X.new_workbook(); ws = wb.active; th = X.theme()
    c1, _ = X.gate_cell(ws, 1, 1, "=規格待補", th=th)
    c2, _ = X.gate_cell(ws, 2, 1, "=SUM(A1:A2)", th=th)
    assert c1.data_type == "s", "文字型 = 開頭應強制字串"
    assert c2.data_type == "f", "真公式應為 formula"


def test_block_table_none_pending():
    """None 值 → 待決議 + 計入 pending。"""
    wb = X.new_workbook(); ws = wb.active; th = X.theme()
    end, npend = X.block_table(ws, 1, 1, ["a", "b"], [[1, None], [None, 2]], th=th)
    assert npend == 2, f"兩個 None 應計 2 pending，實得 {npend}"
    assert ws.cell(2, 2).value == "待決議"  # row2 的 b 欄（第一列資料的 None）


def test_sigma_row_is_live_formula():
    """sigma_row 落的是活公式字串，不是算好的值。"""
    wb = X.new_workbook(); ws = wb.active
    X.sigma_row(ws, 5, 2, 2, 4)
    v = ws.cell(5, 2).value
    assert isinstance(v, str) and v.startswith("=SUM("), f"應為活公式，實得 {v}"


def test_meta_hidden_and_sha():
    """_meta 隱藏 + sha16 戳記。"""
    wb = X.new_workbook()
    src = _tmp(".txt"); src.write_text("hello", encoding="utf-8")
    X.stamp_meta(wb, X.meta_rows({"src": src}))
    assert wb["_meta"].sheet_state == "hidden"
    # 第一列第三欄是 sha16
    sha = wb["_meta"].cell(1, 3).value
    assert sha and sha != "MISSING" and len(sha) == 16
    src.unlink()


# ── ps_config_load ──
def _write_cfg(d: dict) -> pathlib.Path:
    p = _tmp(".json"); p.write_text(json.dumps(d), encoding="utf-8"); return p


def test_config_classify_phase_shape():
    """phase 由表號前綴推、shape 由選擇器推、缺值進 missing。"""
    cfg = _write_cfg({"Bet_Cost": 50, "MainGameReelSetWeight": [30, 40, 30]})
    r = CL.load_config(cfg, None)
    # Bet_Cost 走 board. → General/scalar；reelset → Main/weights
    allitems = [it for items in r["phases"].values() for it in items]
    bet = next(it for it in allitems if it["id"] == "bet_cost")
    assert bet["shape"] == "scalar"
    rw = next(it for it in allitems if it["id"] == "reelset_weight")
    assert rw["value"] == [30, 40, 30]
    assert len(r["missing"]) > 0  # 其他鍵沒給 → 進 missing
    cfg.unlink()


def test_config_table_number():
    """dual[M-2] → table=M-2、phase=Main。"""
    cfg = _write_cfg({"FakeSCScores": [1, 2]})
    r = CL.load_config(cfg, None)
    allitems = [it for items in r["phases"].values() for it in items]
    sc = next(it for it in allitems if it["id"] == "fake_sc_scores")
    assert sc["table"] == "M-2" and sc["value"] == [1, 2]
    cfg.unlink()


# ── ps_xlsx kinds ──
def test_probtable_coverage_and_meta():
    cfg = _write_cfg({"Bet_Cost": 50, "MainGameReelSetWeight": [30, 40, 30]})
    out = pathlib.Path(tempfile.mkdtemp())
    d = PX.build_probtable(cfg, out, "milefo", "彌勒佛", None)
    assert "參數表" in d["sheets"] and "附-1 涵蓋" in d["sheets"]
    wb = openpyxl.load_workbook(d["xlsx"])
    assert wb["_meta"].sheet_state == "hidden"
    cfg.unlink()


def test_versions_detects_diff():
    """兩版權重不同 → diff_rows>=1。"""
    a = _write_cfg({"MainGameReelSetWeight": [30, 40, 30]})
    b = _write_cfg({"MainGameReelSetWeight": [25, 50, 25]})
    out = pathlib.Path(tempfile.mkdtemp())
    d = PX.build_versions([str(a), str(b)], out, None)
    assert d["diff_rows"] >= 1, "權重不同應被抓到"
    a.unlink(); b.unlink()


def test_qtreport_golden_cell_and_formula():
    """golden 放可調格、偏差/判定為活公式。"""
    rpt = _tmp(".json")
    rpt.write_text(json.dumps({"games": {"Total": {"rtp": 0.594, "hitrate": None,
                   "freq": None, "multi": None, "maxmulti": 166.8}}}), encoding="utf-8")
    out = pathlib.Path(tempfile.mkdtemp())
    d = PX.build_qtreport(rpt, out, 0.594)
    wb = openpyxl.load_workbook(d["xlsx"]); ws = wb["快測對照"]
    assert ws["B1"].value == 0.594  # golden 可調格
    # 偏差公式在 G 欄
    gcol = [ws.cell(r, 7).value for r in range(4, 8) if ws.cell(r, 7).value]
    assert any(isinstance(v, str) and v.startswith("=") for v in gcol), "偏差應為活公式"
    rpt.unlink()


def test_check_input_blank_and_threshold():
    """檢核表：輸入格留空（非待決議）、門檻可調格、判定活公式。"""
    spec = _tmp(".yaml")
    spec.write_text(
        'contract: "1"\nname: t\ntitle: T\nthresholds: {warn: 0.03, error: 0.1}\n'
        'columns: [層級, 輸入, 參考, 偏差, 檢核]\n'
        'rows:\n  - [C, null, 100, "=IF(B4=\\"\\",\\"\\",B4/C4-1)", "=IF(B4=\\"\\",\\"未填\\",\\"OK\\")"]\n',
        encoding="utf-8")
    out = pathlib.Path(tempfile.mkdtemp())
    d = PX.build_check(spec, out)
    wb = openpyxl.load_workbook(d["xlsx"]); ws = wb.active
    assert ws["B4"].value is None, "輸入格應留空讓人填"
    assert ws["C1"].value == 0.03, "warn 門檻值應在 C1 可調格"
    assert str(ws["E4"].value).startswith("="), "檢核應為活公式"
    spec.unlink()
