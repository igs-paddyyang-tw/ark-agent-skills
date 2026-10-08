"""ark-game-gdd 守門測試：lint 規則（含真實瑕疵）、build deterministic、戳記 STALE、佔位符與 slot、claim 不多不少。"""
from __future__ import annotations
import csv, json, pathlib, subprocess, sys
import pytest

yaml = pytest.importorskip("yaml")
pytest.importorskip("jinja2")
SCRIPTS = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SCRIPTS.parent)); sys.path.insert(0, str(SCRIPTS / "gdd"))
import gdd_common as C  # noqa: E402

PNG = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0000000020001e221bc330000000049454e44ae426082")


def make_pack(tmp: pathlib.Path, legacy: bool = False, **tweak) -> pathlib.Path:
    """legacy=True：舊中文三夾（圖騰/全示意圖/規格書競品圖）且 gdd.yaml.assets 不寫夾名，驗相容。"""
    D = C.LEGACY_ASSET_DIRS if legacy else C.ASSET_DIRS
    pack = tmp / "demo"; (pack / "rules").mkdir(parents=True)
    for d in D.values():
        (pack / "assets" / d).mkdir(parents=True)
    for f in ("S1.png", "S2.png", "W.png"):
        (pack / "assets" / D["symbols"] / f).write_bytes(PNG)
    for f in ("main_1.png", "bad.png.png", "Snipaste_1.png"):
        (pack / "assets" / D["screens"] / f).write_bytes(PNG)
    (pack / "assets" / D["reference"] / "S1_參考.png").write_bytes(PNG)
    gdd = {"contract": "1", "slug": "demo", "title": "Demo 素材總覽", "short_title": "Demo", "domain": "slot-game", "status": "draft", "distribution": "internal",
           "assets": {"root": "assets"} if legacy else {"root": "assets", **C.ASSET_DIRS},
           "spec": [{"k": "盤面", "short": "3×5", "v": "3×5"}, {"k": "收費", "v": "88"}],
           "symbol_groups": [{"id": "normal", "title": "一般"}, {"id": "special", "title": "特殊"}],
           "odds_order": ["S1", "S2"],
           "features": [{"id": "main", "title": "主遊戲", "rules": "rules/main.md", "screens_title": "主遊戲畫面"}, {"id": "common", "title": "共用素材"}],
           "i18n_columns": [{"id": "en", "label": "英文"}, {"id": "tw", "label": "繁中"}],
           "build": {"output": "out.html"}}
    gdd.update(tweak.get("gdd", {}))
    symbols = [{"code": "S1", "group": "normal", "sym_id": 1, "name": "星", "file": "S1.png", "soft": True, "art_name": "a.png", "ref_files": ["S1_參考.png"], "odds": [50, 20, 10], "scene": "主遊戲"},
               {"code": "S2", "group": "normal", "sym_id": 2, "name": "月", "file": "S2.png", "soft": True, "odds": [30, 10, 5]},
               {"code": "WILD", "group": "special", "sym_id": 3, "name": "百搭", "file": "W.png", "soft": True, "desc": "替代"}]
    symbols = tweak.get("symbols", symbols)
    screens = [{"feature": "main", "step": "1", "title": "待機", "file": "main_1.png", "lang": "英文", "desc": "d"},
               {"feature": "main", "step": "2", "title": "轉場", "file": "bad.png.png", "issue": "顏色對調", "desc": "d"},
               {"feature": "main", "step": "3", "title": "預中", "file": "Snipaste_1.png", "desc": "d"},
               {"feature": "common", "step": "—", "title": "LOGO", "file": None, "from": "美術需求表 1", "desc": "d"}]
    screens = tweak.get("screens", screens)
    info = {"contract": "1", "langs": [{"id": "tw", "label": "繁中"}, {"id": "en", "label": "English"}], "default_lang": "tw",
            "placeholders": {"星圖": "S1"},
            "blocks": [{"t": "h", "en": "ODDS TABLE", "tw": "賠率表", "widget": "odds_grid"},
                       {"t": "p", "en": "When {星圖} appears", "tw": "當{星圖}出現", "id": "b1"}],
            "slots": [{"after": "b1", "items": [{"label": "圖", "file": "main_1.png"}, {"label": "待補圖", "file": None}]}]}
    info = tweak.get("info", info)
    (pack / "gdd.yaml").write_text(yaml.safe_dump(gdd, allow_unicode=True, sort_keys=False), encoding="utf-8")
    (pack / "symbols.yaml").write_text(yaml.safe_dump({"contract": "1", "symbols": symbols}, allow_unicode=True, sort_keys=False), encoding="utf-8")
    (pack / "screens.yaml").write_text(yaml.safe_dump({"contract": "1", "screens": screens}, allow_unicode=True, sort_keys=False), encoding="utf-8")
    (pack / "info.yaml").write_text(yaml.safe_dump(info, allow_unicode=True, sort_keys=False), encoding="utf-8")
    with (pack / "i18n.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(["en", "tw"]); w.writerow(["GRAND", "巨獎"]); w.writerow(["{25} {星圖}", "{25}{星圖}"])
    (pack / "rules" / "main.md").write_text(tweak.get("rules", "### 規則\n\n- **WILD**：替代\n\n| 觸發 | 盤面 |\n|---|---|\n| 藍 | 3×5 |\n"), encoding="utf-8")
    return pack


def run(script: str, *args) -> dict:
    r = subprocess.run([sys.executable, str(SCRIPTS / "gdd" / script), *args], capture_output=True, text=True, encoding="utf-8")
    return json.loads(r.stdout.strip().splitlines()[-1]) | {"rc": r.returncode}


def test_lint_pass_and_catches_real_defects(tmp_path):
    pack = make_pack(tmp_path)
    r = run("gdd_lint.py", "--pack", str(pack))
    assert r["success"] and r["data"]["status"] == "PASS"
    rep = json.loads((pack / "lint-report.json").read_text(encoding="utf-8"))
    rules = [w["rule"] for w in rep["warnings"]]
    msgs = " ".join(w["msg"] for w in rep["warnings"])
    assert "GDD-SCR-HYGIENE" in rules and ".png.png" in msgs      # 雙副檔名
    assert "GDD-SCR-ISSUE" in rules and "顏色對調" in msgs           # 待修
    assert "Snipaste_" in msgs                                      # 泛用檔名
    assert "GDD-INFO-TODO" in rules


def test_lint_errors(tmp_path):
    # 圖檔不存在 / odds 不是三值 / 待補無 from / 佔位符未定義 / slot 錨點不存在
    pack = make_pack(tmp_path,
                     symbols=[{"code": "S1", "group": "normal", "sym_id": 1, "name": "星", "file": "missing.png", "odds": [1, 2]}],
                     screens=[{"feature": "main", "step": "1", "title": "x", "file": None, "desc": "d"}],
                     info={"langs": [{"id": "tw", "label": "繁中"}], "default_lang": "tw", "placeholders": {},
                           "blocks": [{"t": "p", "tw": "{未知圖}"}], "slots": [{"after": "nope", "items": []}]})
    r = run("gdd_lint.py", "--pack", str(pack))
    assert not r["success"] and r["rc"] == 3
    msgs = " ".join(e["msg"] for e in json.loads((pack / "lint-report.json").read_text(encoding="utf-8"))["errors"])
    for needle in ("圖檔不存在", "3 個整數", "必須附 from", "未在 placeholders", "不是任何 block"):
        assert needle in msgs, needle


def test_duplicate_symbol_code_and_sym_id(tmp_path):
    syms = [{"code": "S1", "group": "normal", "sym_id": 1, "name": "a", "file": "S1.png", "odds": [1, 1, 1]},
            {"code": "S1", "group": "normal", "sym_id": 1, "name": "b", "file": "S2.png", "odds": [1, 1, 1]}]
    r = run("gdd_lint.py", "--pack", str(make_pack(tmp_path, symbols=syms)))
    msgs = " ".join(e["msg"] for e in r["data"]["first_errors"])
    assert "sym_id 1" in msgs and "role=symbol" in msgs


def test_build_deterministic_and_stamp(tmp_path):
    pack = make_pack(tmp_path)
    run("gdd_lint.py", "--pack", str(pack))
    a = run("gdd_build.py", "--pack", str(pack)); assert a["success"], a
    html = (pack / "out.html").read_bytes()
    assert run("gdd_build.py", "--pack", str(pack))["success"] and (pack / "out.html").read_bytes() == html
    assert run("gdd_build.py", "--pack", str(pack), "--check")["data"]["status"] == "OK"
    (pack / "rules" / "main.md").write_text("### 規則\n\n- 改了\n", encoding="utf-8")
    c = run("gdd_build.py", "--pack", str(pack), "--check")
    assert not c["success"] and c["data"]["status"] == "STALE"


def test_html_content(tmp_path):
    pack = make_pack(tmp_path)
    run("gdd_lint.py", "--pack", str(pack)); run("gdd_build.py", "--pack", str(pack))
    h = (pack / "out.html").read_text(encoding="utf-8")
    assert h.count('class="sym"') == 3 and h.count('class="scr"') == 3 and 'class="scr todo"' in h
    assert 'src="assets/symbols/S1.png"' in h and "S1_%E5%8F%83%E8%80%83.png" in h        # linked 模式相對路徑
    assert '<img class="inl" src="assets/symbols/S1.png"' in h                              # 佔位符 → 圖騰
    assert 'class="odds-grid"' in h and h.count('class="odds-cell"') == 4   # 2 語系 × 2 圖騰
    assert 'class="slot" data-gi=' in h and 'class="slot empty"' in h
    assert "<b>WILD</b>" in h and "<th>觸發</th>" in h                                    # rules md → html
    assert 'id="infobox-en" hidden' in h and 'id="infobox-tw">' in h
    assert "<!-- gdd-src: gdd.yaml:" in h
    todo = (pack / "todo.md").read_text(encoding="utf-8")
    assert "LOGO" in todo and ".png.png" in todo and "顏色對調" in todo


def test_inject_blocked(tmp_path):
    pack = make_pack(tmp_path, rules="### 規則\n\n- <script>alert(1)</script>\n")
    r = run("gdd_lint.py", "--pack", str(pack))
    assert not r["success"] and any(e["rule"] == "GDD-INJECT" for e in r["data"]["first_errors"])


def test_md_to_html_subset():
    h = C.md_to_html("### T\n\n> note\n\n- a **b** `c`\n\n| x | y |\n|---|---|\n| 1 | 2 |\n\npara")
    assert "<h3>T</h3>" in h and '<p class="note">note</p>' in h and "<li>a <b>b</b> <code>c</code></li>" in h
    assert "<th>x</th>" in h and "<td>2</td>" in h and "<p>para</p>" in h


# ── gdd_extract ────────────────────────────────────────────────────────────

def _wb(tmp: pathlib.Path) -> pathlib.Path:
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.drawing.image import Image as XImage
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "2. 規格玩法"
    ws["B3"], ws["C3"] = "盤面", "3x5"; ws["B4"], ws["C4"] = "對獎方式", "9 LINES"; ws["B5"], ws["C5"] = "收費", "50"
    ws["B6"], ws["C6"] = "Free Game", "Y"; ws["B8"], ws["C8"] = "JP", "N"
    ws["B16"], ws["C16"], ws["D16"], ws["F16"], ws["G16"], ws["H16"] = "名稱", "圖騰", "SymID", "3連線", "4連線", "5連線"
    ws["B18"], ws["C18"], ws["D18"], ws["F18"], ws["G18"], ws["H18"] = "星", "M1", 11, 20, 125, 1000
    ws["B19"], ws["C19"], ws["D19"], ws["F19"], ws["G19"], ws["H19"] = "月", "M2", 12, 15, 75, 400
    ws["A49"] = "玩法敘述"; ws["B51"] = "1) 主遊戲"; ws["C52"] = "一般 Line Game。"
    s2 = wb.create_sheet("6. 圖騰設計")
    for j, h in enumerate(["圖騰", "參考", "名稱", "類型", "面積", "外框", "色調", "說明"], 1):
        s2.cell(row=2, column=j, value=h)
    s2["A3"], s2["C3"], s2["D3"] = "M1", "星", "一般-強"
    png = tmp / "p.png"; png.write_bytes(PNG)
    im = XImage(str(png)); s2.add_image(im, "B3")
    s3 = wb.create_sheet("11. INFO"); s3["A9"] = "賠率"; s3["A10"] = "當{星圖}出現時，可觸發免費遊戲。"; s3["A11"] = "成本 {80}"
    s4 = wb.create_sheet("12. 多國語系表(選)")
    for j, h in enumerate(["編號", "英文", "繁中", "簡中", "日文", "泰文", "印尼文", "越南文", "用途"], 1):
        s4.cell(row=8, column=j, value=h)
    s4["A9"], s4["B9"], s4["C9"], s4["I9"] = 1, "GRAND", "巨獎", "JP面板"
    s5 = wb.create_sheet("藍玩法"); s5["A1"] = "Free Game 藍寶箱"; s5["A3"] = "❖ 有 Free Game Ball 有機會觸發"; s5["A4"] = "Step 1：兌現獎"
    s6 = wb.create_sheet("9. 美術需求表")
    for j, h in enumerate(["編號", "類型", "製作項目說明", "原畫設定"], 1):
        s6.cell(row=6, column=j, value=h)
    s6["A7"], s6["B7"], s6["C7"] = 1, "圖騰", "圖騰 - M1"
    p = tmp / "spec.xlsx"; wb.save(p); return p


def test_extract_classify_and_pack(tmp_path):
    xlsx = _wb(tmp_path); out = tmp_path / "pack"
    r = run("gdd_extract.py", "--xlsx", str(xlsx), "--out", str(out), "--classify")
    roles = {s["name"]: s["role"] for s in r["data"]["sheets"]}
    assert roles["2. 規格玩法"] == "spec" and roles["6. 圖騰設計"] == "symbols" and roles["11. INFO"] == "info"
    assert roles["12. 多國語系表(選)"] == "i18n" and roles["藍玩法"] == "rules" and roles["9. 美術需求表"] == "art_list"
    r = run("gdd_extract.py", "--xlsx", str(xlsx), "--out", str(out)); assert r["success"], r
    c = r["data"]["counts"]
    assert c["spec"] >= 4 and c["odds"] == 2 and c["symbols"] == 2 and c["i18n"] == 1 and c["rules"] == 2
    syms = yaml.safe_load((out / "symbols.yaml").read_text(encoding="utf-8"))["symbols"]
    m1 = next(s for s in syms if s["code"] == "M1")
    assert m1["sym_id"] == 11 and m1["odds"] == [1000, 125, 20] and m1["file"] and (out / "assets" / "symbols" / m1["file"]).exists()
    info = yaml.safe_load((out / "info.yaml").read_text(encoding="utf-8"))
    assert info["placeholders"] == {"星圖": None} or info["placeholders"].get("星圖") in (None, "M1")
    assert (out / "rules" / "blue.md").exists() and "- 有 Free Game Ball" in (out / "rules" / "blue.md").read_text(encoding="utf-8")
    assert (out / "extract-map.yaml").exists() and (out / "extract-report.md").exists()
    # map 手改不被覆蓋
    mp = (out / "extract-map.yaml").read_text(encoding="utf-8").replace("role: rules", "role: unknown")
    (out / "extract-map.yaml").write_text(mp, encoding="utf-8")
    r2 = run("gdd_extract.py", "--xlsx", str(xlsx), "--out", str(out))
    assert r2["data"]["counts"]["rules"] == 1


# ── gdd_from_spec ──────────────────────────────────────────────────────────

def _run_dir(tmp: pathlib.Path) -> pathlib.Path:
    run = tmp / "20260101-slot-game-demo-001"; (run / "frames" / "keyframes").mkdir(parents=True)
    spec = """---
type: "game-spec"
stage: "draft"
contract: "1"
run_id: "20260101-slot-game-demo-001"
domain: "slot-game"
---

# Game Spec Draft

## 01. Game Overview
<!-- section:overview items:C01 -->

### OBSERVED
- `game.title` = Demo Slot — evidence: E001 — confidence: high

### UNKNOWN
- `game.platform` — question: Q001

## 04. Reel Configuration
<!-- section:reel_config items:S01 -->

### OBSERVED
- `reel.columns` = 5 — evidence: E001 — confidence: high
- `reel.rows` = 3 — evidence: E001 — confidence: high

## 07. Wild
<!-- section:wild items:S03 -->

### OBSERVED
- `wild.present` = true — evidence: E002 — confidence: high

## 10. Free Spin
<!-- section:free_spin items:S06 -->

### OBSERVED
- `free_spin.present` = true — evidence: E003 — confidence: high
- `free_spin.count_awarded` = 8 — evidence: E003 — confidence: high

## 21. Evidence
<!-- section:evidence items: special:evidence -->

- E001 …
"""
    (run / "game-spec.draft.md").write_text(spec, encoding="utf-8")
    (run / "entities.json").write_text(json.dumps({"entities": {"symbol": ["symbol_a", "symbol_b"]}, "all_ids": ["symbol_a", "symbol_b"]}), encoding="utf-8")
    (run / "game-analysis.yaml").write_text(yaml.safe_dump({"items": [{"id": "S02", "entities": [{"id": "symbol_a", "entity_type": "symbol", "provenance": "OBSERVED", "appearance": "red gem"}]}]}), encoding="utf-8")
    for i, lab in enumerate(["fallback", "reel_stop", "feature_transition"]):
        (run / "frames" / "keyframes" / f"kf{i}.png").write_bytes(PNG)
    (run / "frames" / "keyframes.json").write_text(json.dumps({"keyframes": [
        {"idx": i, "file": f"frames/keyframes/kf{i}.png", "label": lab, "ts": f"00:00:0{i}.000", "t": i, "extra": {"note": f"n{i}"}}
        for i, lab in enumerate(["fallback", "reel_stop", "feature_transition"])]}), encoding="utf-8")
    return run


def test_from_spec_to_green(tmp_path):
    rd = _run_dir(tmp_path); out = tmp_path / "pack"
    r = run("gdd_from_spec.py", "--run", str(rd), "--out", str(out)); assert r["success"], r
    d = r["data"]
    assert d["features"][:2] == ["main", "fg"] and d["symbols"] == 3 and d["screens"] == 3 and d["open_questions"] == 1
    g = yaml.safe_load((out / "gdd.yaml").read_text(encoding="utf-8"))
    kv = {x["k"]: x["v"] for x in g["spec"]}
    assert kv["盤面"].startswith("3×5") and kv["Free Game"] == "Y" and kv["機種名"] == "Demo Slot"
    assert g["extract"]["claims"]["盤面"] == "reel.rows/reel.columns" and "evidence" in g["extract"]["skipped_sections"]
    fg = (out / "rules" / "fg.md").read_text(encoding="utf-8")
    assert fg.startswith("<!-- spec:free_spin ") and "`free_spin.count_awarded` = 8 — evidence: E003" in fg   # 逐條繼承
    sc = yaml.safe_load((out / "screens.yaml").read_text(encoding="utf-8"))["screens"]
    assert {s["feature"] for s in sc} == {"main"} and all((out / "assets" / "illustrations" / s["file"]).exists() for s in sc)  # 無 feature 章 → 落 main
    # lint：符號無圖依政策降 warn → 整包綠，build 出 HTML 且錨點保留為註解
    rr = run("gdd_run.py", "--pack", str(out)); assert rr["success"] and rr["data"]["lint"]["status"] == "PASS", rr
    html = (out / "素材總覽.html").read_text(encoding="utf-8")
    assert "<!-- spec:free_spin" in html and "待美術" in html and "Q001" in (out / "todo.md").read_text(encoding="utf-8")



# ── 2.1：企劃樣板三夾與交付夾（對齊 data/references/kaiji-gdd-sample）──

def test_default_asset_dirs_are_sample_names(tmp_path):
    assert C.ASSET_DIRS == {"symbols": "symbols", "screens": "illustrations", "reference": "spec-reference-images"}
    p = C.load_pack(make_pack(tmp_path))
    assert C.asset_url(p, "screens", "main_1.png") == "assets/illustrations/main_1.png" and p["assets"]["legacy"] == []


def test_legacy_chinese_dirs_still_build_with_warning(tmp_path):
    pack = make_pack(tmp_path, legacy=True)
    p = C.load_pack(pack)
    assert p["assets"]["symbols"] == "圖騰" and set(p["assets"]["legacy"]) == {"symbols", "screens", "reference"}
    r = run("gdd_lint.py", "--pack", str(pack))
    rep = json.loads((pack / "lint-report.json").read_text(encoding="utf-8"))
    assert sum(1 for w in rep["warnings"] if w["rule"] == "GDD-ASSET-LEGACY") == 3
    assert run("gdd_build.py", "--pack", str(pack))["success"]


def _layout(d: pathlib.Path) -> list[str]:
    return sorted(x.name for x in d.iterdir())


def test_export_matches_sample_layout(tmp_path):
    pack = make_pack(tmp_path)
    out = tmp_path / "bundle"
    r = run("gdd_export.py", "--pack", str(pack), "--out", str(out)); assert r["success"], r
    assert _layout(out) == sorted(["Demo_素材總覽.html", "symbols", "illustrations", "spec-reference-images"])
    assert sorted(f.name for f in (out / "symbols").iterdir()) == ["S1.png", "S2.png", "W.png"]
    assert sorted(f.name for f in (out / "spec-reference-images").iterdir()) == ["S1_參考.png"]
    html = (out / "Demo_素材總覽.html").read_text(encoding="utf-8")
    assert 'src="symbols/S1.png"' in html and "assets/" not in html and "<title>Demo 素材總覽</title>" in html
    m = json.loads((pack / "export-manifest.json").read_text(encoding="utf-8"))
    assert m["broken_links"] == [] and m["extra_top_level"] == [] and m["counts"]["illustrations"] == 3


def test_export_from_legacy_pack_outputs_english_dirs(tmp_path):
    pack = make_pack(tmp_path, legacy=True)
    out = tmp_path / "bundle"
    assert run("gdd_export.py", "--pack", str(pack), "--out", str(out))["success"]
    assert {"symbols", "illustrations", "spec-reference-images"} <= set(_layout(out)) and "圖騰" not in _layout(out)


def test_export_all_assets_and_bundle_name(tmp_path):
    pack = make_pack(tmp_path)
    (pack / "assets" / "illustrations" / "unused.png").write_bytes(PNG)
    g = yaml.safe_load((pack / "gdd.yaml").read_text(encoding="utf-8")); g["build"]["bundle_html"] = "示範_素材總覽.html"
    (pack / "gdd.yaml").write_text(yaml.safe_dump(g, allow_unicode=True, sort_keys=False), encoding="utf-8")
    out = tmp_path / "b1"
    assert run("gdd_export.py", "--pack", str(pack), "--out", str(out))["success"]
    assert not (out / "illustrations" / "unused.png").exists() and (out / "示範_素材總覽.html").exists()
    out2 = tmp_path / "b2"
    assert run("gdd_export.py", "--pack", str(pack), "--out", str(out2), "--all-assets")["success"]
    assert (out2 / "illustrations" / "unused.png").exists()


def test_export_blocks_broken_links(tmp_path):
    pack = make_pack(tmp_path)
    (pack / "assets" / "symbols" / "S2.png").unlink()
    r = run("gdd_export.py", "--pack", str(pack), "--out", str(tmp_path / "b"))
    assert not r["success"] and r["rc"] == 3 and r["error"]["code"] == "GATE_BLOCKED"


@pytest.mark.parametrize("evil", ["../escape.png", "sub/x.png", "..\\x.png", "/etc/passwd", "C:x.png", ".."])
def test_export_blocks_path_traversal_names(tmp_path, evil):
    """反證：檔名含路徑必須在任何複製前擋下（exit 3），交付夾外不得出現檔案。"""
    syms = [{"code": "S1", "group": "normal", "sym_id": 1, "name": "星", "file": "S1.png", "soft": True,
             "ref_files": [evil], "odds": [50, 20, 10]}]
    pack = make_pack(tmp_path, symbols=syms)
    out = tmp_path / "b"
    r = run("gdd_export.py", "--pack", str(pack), "--out", str(out))
    assert not r["success"] and r["rc"] == 3 and r["error"]["code"] == "GATE_BLOCKED"
    assert r["data"]["unsafe_names"][0]["file"] == evil
    assert not out.exists() and not (tmp_path / "escape.png").exists()


def test_gdd_run_export_flag(tmp_path):
    pack = make_pack(tmp_path, screens=[{"feature": "main", "step": "1", "title": "待機", "file": "main_1.png", "desc": "d"}])
    r = run("gdd_run.py", "--pack", str(pack), "--export", "--export-out", str(tmp_path / "b"))
    assert r["success"], r
    assert r["data"]["export"]["html"] == "Demo_素材總覽.html" and (tmp_path / "b" / "symbols" / "W.png").exists()
